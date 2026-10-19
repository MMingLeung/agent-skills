"""work-doc-writer 脚本回归测试：python -m pytest <skill>/tests -q"""
import json
import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from PIL import Image, ImageDraw

SCRIPTS = Path(__file__).resolve().parent.parent / "scripts"
sys.path.insert(0, str(SCRIPTS))
import annotate  # noqa: E402

FRONT = "---\nprofile: manual\ntitle: 测试工具\nversions:\n  - date: 2026.10.19\n    version: V1.0\n---\n"


def run(script, *args, cwd=None):
    return subprocess.run([sys.executable, str(SCRIPTS / script), *map(str, args)], capture_output=True,
                          text=True, encoding="utf-8", errors="replace", cwd=cwd)


def make_png(path, size=(400, 300), color="white"):
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", size, color).save(path)
    return path


def test_check_doc_flags_bad_sample(tmp_path):
    md = tmp_path / "bad.md"
    md.write_text(
        "---\nprofile: manual\ntitle: 坏样例\n---\n\n# 操作说明\n## 打包功能\n### 操作步骤\n"
        "1. 双击”main.exe”启动程序。\n2. 打开EXCECL文件，点击“开始。\n"
        "![图1-2 界面](shots/a.png)\n", encoding="utf-8")
    res = run("check_doc.py", md)
    assert res.returncode == 1
    out = res.stdout
    assert "文件头缺少 versions" in out
    assert "第9行 中文引号不成对" in out
    assert "错别字「EXCECL」" in out
    assert "图片不存在" in out
    assert "图号不连续" in out


def test_check_doc_passes_good_sample(tmp_path):
    make_png(tmp_path / "shots" / "1-1.png")
    md = tmp_path / "good.md"
    md.write_text(FRONT + "# 操作说明\n## 打包功能\n### 准备数据\n1. 准备“模板”Excel 文件。\n"
                  "### 操作步骤\n1. 点击“开始”按钮。\n![图1-1 打包界面图](shots/1-1.png)\n"
                  "### 结果说明\n系统会输出结果文件。\n", encoding="utf-8")
    res = run("check_doc.py", md)
    assert res.returncode == 0, res.stdout
    assert "errors=0" in res.stdout


def test_check_doc_code_dir_lookup(tmp_path):
    code = tmp_path / "src"
    (code / "dist").mkdir(parents=True)
    (code / "gui.py").write_text('btn = QPushButton("开始")\n', encoding="utf-8")
    (code / "dist" / "App.exe").write_bytes(b"MZ")
    md = tmp_path / "doc.md"
    md.write_text(FRONT + "# 操作说明\n## 功能\n### 准备数据\n1. 准备数据。\n### 操作步骤\n"
                  "1. 双击“App.exe”启动程序。\n2. 点击“开始”按钮。\n3. 点击“不存在的按钮”。\n"
                  "### 结果说明\n系统会输出结果。\n", encoding="utf-8")
    res = run("check_doc.py", md, "--code-dir", code)
    assert "「不存在的按钮」在源码里没找到" in res.stdout
    assert "「开始」" not in res.stdout
    assert "「App.exe」" not in res.stdout


def test_place_label_prefers_blank_side():
    im = Image.new("L", (600, 300), 255)
    d = ImageDraw.Draw(im)
    for x in range(0, 600, 4):
        d.line([(x, 0), (x, 95)], fill=0)
    box = (100, 100, 250, 160)
    x, y = annotate.place_label(box, (120, 30), im, [box], "auto", 6)
    assert x >= box[2], f"标签应放在右侧空白处，实际 {x},{y}"


def test_annotate_apply_draws_box_and_copies_plain(tmp_path):
    make_png(tmp_path / "shots" / "a.png")
    make_png(tmp_path / "shots" / "b.png")
    spec = {
        "source_dir": "shots", "output_dir": "out",
        "shots": [
            {"file": "a.png", "id": "1-1", "marks": [{"box": [50, 50, 200, 120], "label": "1、点击开始"}]},
            {"file": "b.png", "id": "1-2"},
        ],
    }
    (tmp_path / "shots.json").write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
    res = run("annotate.py", "apply", tmp_path / "shots.json")
    assert res.returncode == 0, res.stderr
    out = Image.open(tmp_path / "out" / "1-1.png").convert("RGB")
    assert out.getpixel((50, 85)) == annotate.RED
    assert (tmp_path / "out" / "1-2.png").exists()
    table = run("annotate.py", "table", tmp_path / "shots.json").stdout
    assert "| 1 | a.png |" in table and "1、点击开始" in table


def build_sample(tmp_path, image="shots/1-1.png"):
    make_png(tmp_path / "shots" / "1-1.png", size=(1600, 1000))
    md = tmp_path / "manual.md"
    md.write_text(FRONT + "# 摘要 {-}\n不编号标题。\n# 操作说明\n## 打包功能\n### 操作步骤\n"
                  "1. 双击“main.exe”启动程序。\n2. 点击“开始”按钮。\n"
                  f"![图1-1 打包界面图]({image})\n表1 参数说明\n| 参数 | 说明 |\n|---|---|\n| 超时 | 30 秒 |\n",
                  encoding="utf-8")
    return md, run("build_docx.py", md)


def num_id(p_or_style_el):
    ppr = p_or_style_el.find(qn("w:pPr"))
    el = ppr.find(qn("w:numPr")) if ppr is not None else None
    return el.find(qn("w:numId")).get(qn("w:val")) if el is not None else None


def test_build_docx_structure(tmp_path):
    md, res = build_sample(tmp_path)
    assert res.returncode == 0, res.stderr
    doc = Document(tmp_path / "manual.docx")
    styles = doc.styles
    assert num_id(styles["Heading 1"].element) == "90"
    assert styles["Heading 3"].paragraph_format.keep_with_next is True
    assert styles["Heading 4"].paragraph_format.keep_with_next is False
    assert styles["Normal"].paragraph_format.line_spacing_rule == WD_LINE_SPACING.EXACTLY
    assert abs(doc.sections[0].left_margin.cm - 3.17) < 0.01

    paras = doc.paragraphs
    unnumbered = next(p for p in paras if p.text == "摘要")
    assert num_id(unnumbered._p) == "0"
    steps = [p for p in paras if p.style.name == "Heading 4"]
    assert [p.text for p in steps] == ["双击“main.exe”启动程序。", "点击“开始”按钮。"]

    img_idx = next(i for i, p in enumerate(paras) if p._p.findall(".//" + qn("w:drawing")))
    img_p = paras[img_idx]
    assert img_p.paragraph_format.line_spacing_rule != WD_LINE_SPACING.EXACTLY
    assert paras[img_idx + 1].text == "图1-1 打包界面图"
    assert paras[img_idx + 1].style.name == "Caption"
    width_cm = img_p._p.find(".//" + qn("wp:extent")).get("cx")
    assert int(width_cm) / 360000 <= 14.5 + 0.01

    assert len(doc.tables) == 2
    assert doc.tables[1].cell(1, 1).text == "30 秒"
    assert any(p.text == "表1 参数说明" and p.style.name == "Caption" for p in paras)


def test_build_docx_missing_image_fails(tmp_path):
    md, res = build_sample(tmp_path, image="shots/missing.png")
    assert res.returncode == 1
    assert "图片不存在" in res.stderr
    assert not (tmp_path / "manual.docx").exists()


@pytest.mark.parametrize("name", ["manual.md", "spec.md", "report.md"])
def test_templates_have_valid_front_matter(name):
    import yaml
    text = (SCRIPTS.parent / "templates" / name).read_text(encoding="utf-8")
    meta = yaml.safe_load(text.split("---", 2)[1])
    assert meta["profile"] in {"manual", "spec", "report"}
    assert meta["title"]
