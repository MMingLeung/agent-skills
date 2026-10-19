"""Markdown 源稿 → 公文体 Word（可选更新目录、导出 PDF、渲染预览图）。

用法：
  python build_docx.py <源稿.md> [-o 输出.docx] [--finalize] [--pdf] [--preview 预览目录]

源稿语法见 reference/layout.md。
"""
import argparse
import re
import subprocess
import sys
from pathlib import Path

import yaml
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

HERE = Path(__file__).resolve().parent

PROFILES = {
    "manual": {"margins": (2.54, 2.54, 3.17, 3.17), "cover": True, "toc": True, "page_number": False, "img_w": 14.5, "img_h": 20.0},
    "spec": {"margins": (2.54, 2.54, 3.17, 3.17), "cover": True, "toc": True, "page_number": False, "img_w": 14.5, "img_h": 20.0},
    "report": {"margins": (2.6, 2.6, 2.8, 2.6), "cover": False, "toc": False, "page_number": True, "img_w": 15.5, "img_h": 21.0},
}
DOC_TYPE_DEFAULT = {"manual": "操作手册", "spec": "说明书", "report": ""}

BODY_FONT = "仿宋_GB2312"
TITLE_FONT = "方正小标宋简体"
HEADING_FONTS = {1: ("黑体", False), 2: ("楷体", False), 3: ("仿宋", True), 4: ("仿宋", False)}
NUM_FORMATS = [("chineseCounting", "%1、"), ("chineseCounting", "（%2）"), ("decimal", "%3．"), ("decimal", "（%4）")]
SZ_BODY, SZ_CAPTION, SZ_TITLE, SZ_TABLE, SZ_TOC = Pt(16), Pt(10.5), Pt(22), Pt(12), Pt(14)
LINE_EXACT = Pt(28)
INDENT_2 = Pt(32)
NUM_ID = 90

RE_HEADING = re.compile(r"^(#{1,4})\s+(.*?)\s*(\{-\})?\s*$")
RE_ORDERED = re.compile(r"^\d+\.\s+(.*)$")
RE_IMAGE = re.compile(r"^!\[(.*?)\]\((.*?)\)\s*$")
RE_TABLE_CAPTION = re.compile(r"^表\d+(-\d+)?\s")
RE_BOLD = re.compile(r"\*\*(.+?)\*\*")


def set_fonts(rpr_owner, name):
    rpr = rpr_owner.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rfonts.attrib.pop(qn(f"w:{attr}"), None)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)


def style_font(style, name, size, bold=False):
    style.font.size = size
    style.font.bold = bold
    style.font.italic = False
    color = style.element.get_or_add_rPr().find(qn("w:color"))
    if color is not None:
        color.getparent().remove(color)
    style.font.color.rgb = RGBColor(0, 0, 0)
    set_fonts(style.element, name)


def run_font(run, name, size, bold=None):
    run.font.size = size
    if bold is not None:
        run.font.bold = bold
    set_fonts(run._r, name)


def exact_spacing(pf, before=0, after=0):
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = LINE_EXACT
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)


def single_spacing(pf, before=0, after=0):
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)


def add_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(NUM_ID))
    mlt = OxmlElement("w:multiLevelType")
    mlt.set(qn("w:val"), "multilevel")
    abstract.append(mlt)
    for i, (fmt, text) in enumerate(NUM_FORMATS):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(i))
        for tag, val in (("w:start", "1"), ("w:numFmt", fmt), ("w:pStyle", f"Heading{i + 1}"),
                         ("w:suff", "nothing"), ("w:lvlText", text), ("w:lvlJc", "left")):
            el = OxmlElement(tag)
            el.set(qn("w:val"), val)
            lvl.append(el)
        abstract.append(lvl)
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(NUM_ID))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(NUM_ID))
    num.append(ref)
    numbering.append(num)


def set_num(ppr, ilvl, num_id):
    numpr = ppr.get_or_add_numPr()
    numpr.get_or_add_ilvl().val = ilvl
    numpr.get_or_add_numId().val = num_id


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    style_font(normal, BODY_FONT, SZ_BODY)
    exact_spacing(normal.paragraph_format)
    normal.paragraph_format.first_line_indent = INDENT_2
    add_numbering(doc)
    for lvl, (font, bold) in HEADING_FONTS.items():
        s = styles[f"Heading {lvl}"]
        style_font(s, font, SZ_BODY, bold)
        pf = s.paragraph_format
        exact_spacing(pf, before=6 if lvl == 1 else 0)
        pf.first_line_indent = INDENT_2 if lvl >= 3 else Pt(0)
        pf.left_indent = Pt(0)
        # 四级标题承载操作步骤，属于正文；若也与下段同页，整串步骤会连成一块被推到下一页
        pf.keep_with_next = lvl < 4
        set_num(s.element.get_or_add_pPr(), lvl - 1, NUM_ID)
    try:
        cap = styles["Caption"]
    except KeyError:
        cap = styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    style_font(cap, BODY_FONT, SZ_CAPTION)
    single_spacing(cap.paragraph_format, after=6)
    cap.paragraph_format.first_line_indent = Pt(0)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i in (1, 2, 3):
        try:
            toc = styles[f"toc {i}"]
        except KeyError:
            toc = styles.add_style(f"toc {i}", WD_STYLE_TYPE.PARAGRAPH)
        toc.base_style = normal
        style_font(toc, BODY_FONT, SZ_TOC)
        single_spacing(toc.paragraph_format, after=4)
        toc.paragraph_format.first_line_indent = Pt(0)
        toc.paragraph_format.left_indent = Pt(14 * 2 * (i - 1))


def no_indent(p):
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.left_indent = Pt(0)


def add_inline(p, text, size=None, font=None):
    pos = 0
    for m in RE_BOLD.finditer(text):
        if m.start() > pos:
            r = p.add_run(text[pos:m.start()])
            if font:
                run_font(r, font, size)
        r = p.add_run(m.group(1))
        r.bold = True
        if font:
            run_font(r, font, size)
        pos = m.end()
    if pos < len(text):
        r = p.add_run(text[pos:])
        if font:
            run_font(r, font, size)


def add_field(p, instr, placeholder=""):
    def el(tag, **attrs):
        e = OxmlElement(tag)
        for k, v in attrs.items():
            e.set(qn(k), v)
        return e
    r = p.add_run()
    r._r.append(el("w:fldChar", **{"w:fldCharType": "begin"}))
    r = p.add_run()
    it = el("w:instrText", **{"xml:space": "preserve"})
    it.text = instr
    r._r.append(it)
    r = p.add_run()
    r._r.append(el("w:fldChar", **{"w:fldCharType": "separate"}))
    p.add_run(placeholder)
    r = p.add_run()
    r._r.append(el("w:fldChar", **{"w:fldCharType": "end"}))


def page_break(doc):
    p = doc.add_paragraph()
    no_indent(p)
    p.add_run().add_break(WD_BREAK.PAGE)


def add_title_lines(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        no_indent(p)
        single_spacing(p.paragraph_format, after=6)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_font(p.add_run(line), TITLE_FONT, SZ_TITLE)


def add_cover(doc, meta, profile):
    for _ in range(5):
        no_indent(doc.add_paragraph())
    title = meta["title"]
    lines = title if isinstance(title, list) else [title]
    doc_type = meta.get("doc_type", DOC_TYPE_DEFAULT[profile])
    add_title_lines(doc, lines + ([doc_type] if doc_type else []))
    for _ in range(3):
        no_indent(doc.add_paragraph())
    versions = meta.get("versions") or []
    if versions:
        add_table(doc, [["日期", "版本"]] + [[str(v["date"]), str(v["version"])] for v in versions], width_cm=8)
    page_break(doc)


def add_toc(doc):
    p = doc.add_paragraph()
    no_indent(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_font(p.add_run("目  录"), "黑体", SZ_BODY)
    p = doc.add_paragraph()
    no_indent(p)
    add_field(p, 'TOC \\o "1-3" \\h \\z \\u', "（目录未更新：在 Word 中右键此处 → 更新域）")
    page_break(doc)


def add_heading(doc, level, text, numbered=True):
    p = doc.add_paragraph(style=f"Heading {level}")
    add_inline(p, text)
    if not numbered:
        set_num(p._p.get_or_add_pPr(), 0, 0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)


def add_paragraph(doc, text):
    p = doc.add_paragraph()
    add_inline(p, text)


def add_image(doc, path, caption, prof):
    with Image.open(path) as im:
        w, h = im.size
    natural_cm = w / 96 * 2.54
    width_cm = min(prof["img_w"], natural_cm)
    if width_cm * h / w > prof["img_h"]:
        width_cm = prof["img_h"] * w / h
    p = doc.add_paragraph()
    no_indent(p)
    single_spacing(p.paragraph_format, before=6)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = bool(caption)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    if caption:
        doc.add_paragraph(caption, style="Caption")


def add_table(doc, rows, caption=None, width_cm=None):
    if caption:
        cp = doc.add_paragraph(caption, style="Caption")
        cp.paragraph_format.keep_with_next = True
        cp.paragraph_format.space_before = Pt(6)
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(ncol):
            cell = t.cell(i, j)
            if width_cm:
                cell.width = Cm(width_cm / ncol)
            p = cell.paragraphs[0]
            no_indent(p)
            single_spacing(p.paragraph_format, before=2, after=2)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            text = row[j] if j < len(row) else ""
            add_inline(p, text, SZ_TABLE, BODY_FONT)
            if i == 0:
                for r in p.runs:
                    r.bold = True
    return t


def add_page_number(doc):
    for section in doc.sections:
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        no_indent(p)
        run_font(p.add_run("— "), "宋体", Pt(14))
        add_field(p, "PAGE", "1")
        run_font(p.add_run(" —"), "宋体", Pt(14))


def parse_source(text):
    if text.startswith("---"):
        _, fm, body = text.split("---", 2)
        return yaml.safe_load(fm) or {}, body
    return {}, text


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def render_body(doc, body, base_dir, prof, errors):
    lines = body.splitlines()
    i = 0
    pending_caption = None
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        i += 1
        if not stripped:
            continue
        if stripped == "<!-- pagebreak -->":
            page_break(doc)
            continue
        if stripped.startswith("<!--"):
            continue
        if stripped.startswith("|"):
            rows = [split_row(stripped)]
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            rows = [r for r in rows if not all(re.fullmatch(r":?-{2,}:?", c) for c in r if c)]
            add_table(doc, rows, caption=pending_caption)
            pending_caption = None
            continue
        if pending_caption:
            add_paragraph(doc, pending_caption)
            pending_caption = None
        m = RE_HEADING.match(stripped)
        if m:
            add_heading(doc, len(m.group(1)), m.group(2), numbered=not m.group(3))
            continue
        m = RE_IMAGE.match(stripped)
        if m:
            img = Path(m.group(2))
            img = img if img.is_absolute() else base_dir / img
            if not img.exists():
                errors.append(f"图片不存在：{img}")
                continue
            add_image(doc, img, m.group(1).strip(), prof)
            continue
        m = RE_ORDERED.match(stripped)
        if m:
            add_heading(doc, 4, m.group(1))
            continue
        if RE_TABLE_CAPTION.match(stripped) and i < len(lines) and lines[i].strip().startswith("|"):
            pending_caption = stripped
            continue
        if stripped.startswith("- "):
            stripped = stripped[2:]
        add_paragraph(doc, stripped)
    if pending_caption:
        add_paragraph(doc, pending_caption)


def build(src, out):
    meta, body = parse_source(src.read_text(encoding="utf-8"))
    profile = meta.get("profile", "manual")
    if profile not in PROFILES:
        sys.exit(f"profile 只能是 {list(PROFILES)}，当前：{profile}")
    if not meta.get("title"):
        sys.exit("源稿缺少 title")
    prof = PROFILES[profile]
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    top, bottom, left, right = prof["margins"]
    sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin = Cm(top), Cm(bottom), Cm(left), Cm(right)
    setup_styles(doc)
    if prof["cover"]:
        add_cover(doc, meta, profile)
    else:
        title = meta["title"]
        add_title_lines(doc, title if isinstance(title, list) else [title])
        no_indent(doc.add_paragraph())
    if prof["toc"] and meta.get("toc", True):
        add_toc(doc)
    errors = []
    render_body(doc, body, src.parent, prof, errors)
    if prof["page_number"]:
        add_page_number(doc)
    if errors:
        for e in errors:
            print("ERROR", e, file=sys.stderr)
        sys.exit(1)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(f"DOCX {out}")


def finalize(docx, pdf=None):
    cmd = ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(HERE / "office_finalize.ps1"),
           "-Docx", str(docx)]
    if pdf:
        cmd += ["-Pdf", str(pdf)]
    res = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace")
    print(res.stdout.strip())
    if res.returncode != 0:
        print(res.stderr.strip(), file=sys.stderr)
        return False
    return True


def render_preview(pdf, out_dir):
    import fitz
    out_dir.mkdir(parents=True, exist_ok=True)
    with fitz.open(pdf) as doc:
        print(f"PAGES {doc.page_count}")
        for i, page in enumerate(doc, 1):
            path = out_dir / f"page-{i:02d}.png"
            page.get_pixmap(matrix=fitz.Matrix(1.2, 1.2)).save(str(path))
            print(f"PREVIEW {path}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("source")
    ap.add_argument("-o", "--output")
    ap.add_argument("--finalize", action="store_true", help="用 Word/WPS 更新目录域并保存")
    ap.add_argument("--pdf", action="store_true", help="同时导出 PDF（隐含 --finalize）")
    ap.add_argument("--preview", help="渲染每页 PNG 到此目录（隐含 --finalize）")
    args = ap.parse_args()
    src = Path(args.source).resolve()
    out = Path(args.output).resolve() if args.output else src.with_suffix(".docx")
    build(src, out)
    if not (args.finalize or args.pdf or args.preview):
        return
    pdf = out.with_suffix(".pdf") if args.pdf else None
    if args.preview and not pdf:
        pdf = Path(args.preview).resolve() / f"{out.stem}.preview.pdf"
        pdf.parent.mkdir(parents=True, exist_ok=True)
    if not finalize(out, pdf):
        sys.exit("FINALIZE_FAILED：目录未更新，需在 Word 中右键目录 → 更新域")
    if pdf:
        print(f"PDF {pdf}")
    if args.preview:
        render_preview(pdf, Path(args.preview).resolve())


if __name__ == "__main__":
    main()
